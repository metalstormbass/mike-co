"""
Document Parsers
Extract text content from various document formats
"""

from abc import ABC, abstractmethod
from typing import List, Dict, Any, Optional
import io


class BaseParser(ABC):
    """Base class for document parsers"""
    
    @abstractmethod
    async def parse(self, content: bytes, filename: str) -> Dict[str, Any]:
        """
        Parse document content
        
        Returns:
            {
                "text": str,           # Extracted text content
                "metadata": dict,      # Document metadata
                "pages": list[str],    # Optional: text per page
            }
        """
        pass


class TextParser(BaseParser):
    """Parser for plain text files (.txt, .md)"""
    
    async def parse(self, content: bytes, filename: str) -> Dict[str, Any]:
        # Try different encodings
        for encoding in ['utf-8', 'latin-1', 'cp1252']:
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = content.decode('utf-8', errors='replace')
        
        return {
            "text": text,
            "metadata": {
                "encoding": encoding,
                "char_count": len(text),
                "line_count": len(text.split('\n'))
            },
            "pages": [text]  # Single page for text files
        }


class PDFParser(BaseParser):
    """Parser for PDF files"""
    
    async def parse(self, content: bytes, filename: str) -> Dict[str, Any]:
        from PyPDF2 import PdfReader
        
        reader = PdfReader(io.BytesIO(content))
        pages = []
        
        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)
        
        full_text = "\n\n".join(pages)
        
        # Extract metadata
        metadata = {
            "page_count": len(reader.pages),
            "char_count": len(full_text),
        }
        
        if reader.metadata:
            if reader.metadata.title:
                metadata["title"] = reader.metadata.title
            if reader.metadata.author:
                metadata["author"] = reader.metadata.author
            if reader.metadata.subject:
                metadata["subject"] = reader.metadata.subject
        
        return {
            "text": full_text,
            "metadata": metadata,
            "pages": pages
        }


class DocxParser(BaseParser):
    """Parser for Word documents (.docx)"""
    
    async def parse(self, content: bytes, filename: str) -> Dict[str, Any]:
        from docx import Document
        
        doc = Document(io.BytesIO(content))
        
        # Extract paragraphs
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        
        # Extract tables
        tables_text = []
        for table in doc.tables:
            table_content = []
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                table_content.append(" | ".join(row_text))
            tables_text.append("\n".join(table_content))
        
        full_text = "\n\n".join(paragraphs)
        if tables_text:
            full_text += "\n\n[Tables]\n" + "\n\n".join(tables_text)
        
        metadata = {
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
            "char_count": len(full_text)
        }
        
        # Try to get core properties
        try:
            if doc.core_properties.title:
                metadata["title"] = doc.core_properties.title
            if doc.core_properties.author:
                metadata["author"] = doc.core_properties.author
        except:
            pass
        
        return {
            "text": full_text,
            "metadata": metadata,
            "pages": [full_text]  # DOCX doesn't have clear page boundaries
        }


class HTMLParser(BaseParser):
    """Parser for HTML files"""
    
    async def parse(self, content: bytes, filename: str) -> Dict[str, Any]:
        from bs4 import BeautifulSoup
        
        # Parse HTML
        soup = BeautifulSoup(content, 'html.parser')
        
        # Remove script and style elements
        for element in soup(['script', 'style', 'nav', 'footer', 'header']):
            element.decompose()
        
        # Get text
        text = soup.get_text(separator='\n', strip=True)
        
        # Get title if present
        metadata = {
            "char_count": len(text)
        }
        
        title_tag = soup.find('title')
        if title_tag:
            metadata["title"] = title_tag.get_text()
        
        return {
            "text": text,
            "metadata": metadata,
            "pages": [text]
        }


# Parser factory
PARSERS = {
    ".txt": TextParser(),
    ".md": TextParser(),
    ".markdown": TextParser(),
    ".pdf": PDFParser(),
    ".docx": DocxParser(),
    ".html": HTMLParser(),
    ".htm": HTMLParser(),
}


def get_parser(filename: str) -> Optional[BaseParser]:
    """Get appropriate parser for file type"""
    import os
    ext = os.path.splitext(filename.lower())[1]
    return PARSERS.get(ext)


def supported_extensions() -> List[str]:
    """Get list of supported file extensions"""
    return list(PARSERS.keys())
